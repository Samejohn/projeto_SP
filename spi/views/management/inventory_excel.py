import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from spi.models import Inventario


@login_required
def exportar_inventario_excel(request):
    # 1. Recupera o parâmetro de busca da URL
    search_query = request.GET.get('q', '').strip()

    # 2. Filtra os dados conforme a busca usando Q objects
    queryset = Inventario.objects.all()
    if search_query:
        queryset = queryset.filter(
            Q(numero_patrimonio__icontains=search_query) |
            Q(item_modelo__icontains=search_query) |
            Q(setor__icontains=search_query)
        ).distinct()

    # 3. Cria a pasta de trabalho Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventário"

    # 4. Estilos do Excel
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    right_alignment = Alignment(horizontal="right", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # 5. Define os cabeçalhos das colunas
    headers = [
        "Nº Patrimônio", "ID Ativo", "Categoria", "Modelo", 
        "Nº(Série)Licença", "Aquisição", "Quant.", "Valor (R$)", 
        "Total (R$)", "Garantia", "Status", "Setor", "Usuário", "Observações"
    ]
    ws.append(headers)

    # Aplica estilos ao cabeçalho
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment

    # 6. Adiciona os dados das linhas
    for item in queryset:
        row = [
            str(item.numero_patrimonio or ""),
            str(item.id_ativo or ""),
            item.get_categoria_display() if hasattr(item, 'get_categoria_display') else str(item.categoria or ""),
            str(item.item_modelo or ""),
            str(item.serie_licenca or "-"),
            item.data_aquisicao.strftime("%d/%m/%Y") if item.data_aquisicao else "-",
            item.quantidade or 0,
            float(item.valor or 0),
            float(item.total or 0),
            item.validade_garantia.strftime("%d/%m/%Y") if item.validade_garantia else "-",
            item.get_status_display() if hasattr(item, 'get_status_display') else str(item.status or ""),
            str(item.setor or ""),
            str(item.usuario or "-"),
            str(item.observacoes or "-")
        ]
        ws.append(row)

    # 7. Formatação de células, alinhamento e bordas (executa somente se houver linhas de dados)
    if ws.max_row >= 2:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for col_idx, cell in enumerate(row, 1):
                cell.border = thin_border
                
                if col_idx in [6, 10]:  # Datas
                    cell.alignment = center_alignment
                elif col_idx in [7]:    # Quantidade
                    cell.alignment = right_alignment
                    cell.number_format = '#,##0'
                elif col_idx in [8, 9]: # Valores monetários
                    cell.alignment = right_alignment
                    cell.number_format = 'R$ #,##0.00'
                elif col_idx in [1, 2]: # Patrimônio e ID Ativo
                    cell.alignment = center_alignment
                else:
                    cell.alignment = left_alignment

    # 8. Ajusta a largura das colunas automaticamente
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 9. Prepara e RETORNA a resposta HTTP obrigatória
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="inventario_patrimonio.xlsx"'
    wb.save(response)
    
    return response